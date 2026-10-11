"""Lectura de por_radicar.xlsx y escritura de resultados.xlsx."""
import logging
import re
import unicodedata
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from reglas import Fila, texto

log = logging.getLogger("robot")

COLUMNAS_ENTRADA = ["IdLista", "Factura", "NIT", "Proveedor", "Valor", "CentroCosto", "FechaFactura"]
COLUMNAS_SALIDA = ["IdLista", "Factura", "NIT", "Estado", "Causacion", "Detalle"]


def _norm(encabezado) -> str:
    s = unicodedata.normalize("NFKD", str(encabezado or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", "", s).lower()


def _id_lista(valor):
    """Entero o None si la celda no trae un IdLista utilizable."""
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, int):
        return valor
    if isinstance(valor, float):
        return int(valor) if valor.is_integer() else None
    s = str(valor).strip()
    if re.fullmatch(r"\d+(\.0+)?", s):
        return int(float(s))
    return None


def leer_cola(ruta: Path) -> list[Fila]:
    """Primera hoja del libro, según los encabezados; ignora filas sin IdLista."""
    wb = load_workbook(ruta, data_only=True)
    ws = wb.worksheets[0]
    filas = ws.iter_rows(values_only=True)
    encabezados = next(filas, None)
    if encabezados is None:
        return []
    pos = {_norm(h): i for i, h in enumerate(encabezados) if h is not None}
    faltan = [c for c in COLUMNAS_ENTRADA if _norm(c) not in pos]
    if faltan:
        raise ValueError(f"Faltan columnas en {ruta.name}: {', '.join(faltan)}")

    def celda(fila, nombre):
        i = pos[_norm(nombre)]
        return fila[i] if i < len(fila) else None

    cola = []
    for n, fila in enumerate(filas, start=2):
        if all(v is None or str(v).strip() == "" for v in fila):
            continue
        id_lista = _id_lista(celda(fila, "IdLista"))
        if id_lista is None:
            log.warning("Fila %d ignorada: no tiene IdLista", n)
            continue
        cola.append(Fila(
            id_lista=id_lista,
            factura=texto(celda(fila, "Factura")),
            nit_crudo=celda(fila, "NIT"),
            proveedor=texto(celda(fila, "Proveedor")),
            valor_crudo=celda(fila, "Valor"),
            centro=texto(celda(fila, "CentroCosto")),
            fecha_cruda=celda(fila, "FechaFactura"),
        ))
    return cola


def escribir_resultados(ruta: Path, resultados: list[dict]):
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultados"
    ws.append(COLUMNAS_SALIDA)
    for r in resultados:
        ws.append([r["IdLista"], r["Factura"], r["NIT"], r["Estado"],
                   r["Causacion"] or None, r["Detalle"] or None])
    tabla = Table(displayName="Resultados", ref=f"A1:F{max(2, len(resultados) + 1)}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tabla)
    for col, ancho in zip("ABCDEF", (10, 14, 14, 16, 18, 70)):
        ws.column_dimensions[col].width = ancho
    wb.save(ruta)
